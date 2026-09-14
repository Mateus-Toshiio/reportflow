from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from datetime import datetime

def converter_data(data):
    try:
        data_convertida = datetime.strptime(data, "%d/%m/%Y")
        return data_convertida
    except:
        return ""

def gerar_excel(servidores, arquivo_saida):
    workbook = Workbook()
    planilha = workbook.active

    planilha.title = 'Relação de Contratos'

    cabecalho = [
        "Nº",
        "MATRÍCULA",
        "NOME",
        "DATA DE NASCIMENTO",
        "CARGO",
        "SETOR",
        "DATA DE ADMISSÃO"
    ]

    planilha.append(cabecalho)

    for celula in planilha[1]:

        celula.font = Font(bold=True)

        celula.fill = PatternFill(
        fill_type="solid",
        fgColor="1F4E78"
        )

    planilha.column_dimensions["A"].width = 6
    planilha.column_dimensions["B"].width = 15
    planilha.column_dimensions["C"].width = 35
    planilha.column_dimensions["D"].width = 22
    planilha.column_dimensions["E"].width = 25
    planilha.column_dimensions["F"].width = 50
    planilha.column_dimensions["G"].width = 22

    planilha.auto_filter.ref = planilha.dimensions

    planilha.freeze_panes = "A2"

    for indice, servidor in enumerate(servidores, start=1):
        planilha.append([
            indice,
            servidor['matricula'],
            servidor['nome'],
            converter_data(servidor['data_nascimento']),
            servidor['cargo'],
            servidor['setor'],
            converter_data(servidor['data_admissao']),
        ])
        
        planilha[f"D{indice + 1}"].number_format = "DD/MM/YYYY"
        planilha[f"G{indice + 1}"].number_format = "DD/MM/YYYY"

    workbook.save(arquivo_saida)